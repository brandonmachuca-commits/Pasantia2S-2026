# process_calendario.py

from pathlib import Path
from docx import Document
import re
from datetime import date

current_year = date.today().year

INPUT_FILE = Path(__file__).parent.parent.parent / "data" / "raw" / "calendario" / "calendario.docx"

OUTPUT_FILE = Path(__file__).parent.parent.parent / "data"/ "processed" / "calendario" / "contenido.txt"

OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)


def limpiar_texto(texto: str) -> str:

    # espacios múltiples
    texto = re.sub(r"[ \t]+", " ", texto)

    # saltos múltiples
    texto = re.sub(r"\n\s*\n+", "\n\n", texto)

    # eliminar basura común
    basura = [
        "Microsoft Word",
        "Normal",
        "Title",
        "Subtitle",
    ]

    for b in basura:
        texto = texto.replace(b, "")

    return texto.strip()


def leer_docx(path: Path) -> str:

    print(f"[INFO] Leyendo DOCX: {path.name}")

    doc = Document(path)

    lineas = []

    # párrafos normales
    for p in doc.paragraphs:

        texto = p.text.strip()

        if texto:
            lineas.append(texto)

    # tablas
    for table in doc.tables:

        for row in table.rows:

            fila = []

            for cell in row.cells:

                texto = cell.text.strip()

                if texto:
                    fila.append(texto)

            if fila:
                lineas.append(" | ".join(fila))

    return "\n".join(lineas)


def estructurar_calendario(texto: str) -> str:

    # mejorar legibilidad
    reemplazos = {
        "Exámenes Febrero": "\nEXÁMENES FEBRERO",
        "Primer semestre": "\nPRIMER SEMESTRE",
        "Parcial 1": "\nPARCIAL 1",
        "Parcial 2": "\nPARCIAL 2",
        "Exámenes Julio": "\nEXÁMENES JULIO",
        "Semestre 2": "\nSEGUNDO SEMESTRE",
        "Exámenes Diciembre": "\nEXÁMENES DICIEMBRE",
        "Feriados": "\nFERIADOS",
    }

    for viejo, nuevo in reemplazos.items():
        texto = texto.replace(viejo, nuevo)

    return texto


def main():

    if not INPUT_FILE.exists():

        print("[ERROR] No existe:")
        print(INPUT_FILE.resolve())

        return

    texto = leer_docx(INPUT_FILE)

    texto = limpiar_texto(texto)

    texto = estructurar_calendario(texto)

    contenido_final = f"""
================================================================================
CALENDARIO ACADÉMICO {current_year}
================================================================================

{texto}
"""

    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:

        f.write(contenido_final)

    print("\n[OK] contenido.txt generado:")
    print(OUTPUT_FILE.resolve())


if __name__ == "__main__":
    main()