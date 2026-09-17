from pathlib import Path
import re
from openpyxl import load_workbook
from datetime import date

current_year = date.today().year

BASE_DIR = Path(__file__).resolve().parent.parent.parent

INPUT_FILE = (
    BASE_DIR
    / "data"
    / "raw"
    / "docentes"
    / "docentes.xlsx"
)

OUTPUT_FILE = (
    BASE_DIR
    / "data"
    / "processed"
    / "docentes"
    / "contenido.txt"
)

OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)

def limpiar(texto: str) -> str:
    if texto is None:
        return ""
    texto = str(texto)
    texto = texto.replace("\n", " ").replace("\t", " ")
    texto = re.sub(r"\s+", " ", texto)
    return texto.strip()

def separar_docente_email(texto: str):
    match = re.search(r"([a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+)", texto)

    email = ""
    if match:
        email = match.group(1)
        texto = texto.replace(email, "").strip()

    return texto, email

def main():

    print(f"[INFO] Leyendo: {INPUT_FILE}")

    wb = load_workbook(INPUT_FILE, data_only=True)
    ws = wb.active

    lineas = []

    lineas.append("=" * 80)
    lineas.append(f"DOCENTES {current_year}")
    lineas.append("=" * 80)
    lineas.append("")

    docente = ""
    email = ""
    area = ""

    dias = []
    leyendo_tabla = False

    for row in ws.iter_rows(values_only=True):

        fila = [limpiar(c) for c in row if c is not None]
        texto = " ".join(fila).strip()

        if not texto:
            continue

        # -----------------------
        # DOCENTE
        # -----------------------
        if "Docente:" in texto:

            raw = texto.split("Docente:")[-1].strip()

            docente, email = separar_docente_email(raw)

            lineas.append("")
            lineas.append(f"Docente: {docente}")
            lineas.append("-" * 50)

            if email:
                lineas.append(f"Email: {email}")
            leyendo_tabla = False
            continue

        # -----------------------
        # ÁREA
        # -----------------------
        if "Área:" in texto or "Area:" in texto:
            area = texto.replace("Área:", "").replace("Area:", "").strip()
            lineas.append(f"Área: {area}")
            continue

        # -----------------------
        # HEADER TABLA DÍAS
        # -----------------------
        if any(d in texto for d in ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes"]):

            dias = fila
            leyendo_tabla = True
            continue

        # -----------------------
        # FILAS HORARIO
        # -----------------------
        if leyendo_tabla:

            tipo = fila[0] if len(fila) > 0 else ""
            valores = fila[1:]

            if tipo in ["Virtual", "Presencial"]:

                lineas.append(f"- {tipo}")

                for i, val in enumerate(valores):

                    if i < len(dias) - 1:
                        dia = dias[i + 1]
                        if val:
                            lineas.append(f"   {dia}: {val}")

    contenido = "\n".join(lineas)

    with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
        f.write(contenido)

    print("[OK] Archivo generado:")
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()